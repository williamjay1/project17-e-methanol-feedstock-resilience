"""Public-price causal daily MPC and a service-viability tail extension.

Each policy receives current state, current supplier availability and already
published current-day prices. Future prices are past-day profile forecasts
initialized with a public 2021 profile; repair time is never a policy input.
The guard receives
an ex-ante 21-day engineering interruption envelope, not its realization.
"""
import json,time,hashlib
from dataclasses import asdict,replace
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from mucci_service_milp import Design,solve,ROOT,PARAM

def price_forecast():
    w=load_workbook(PARAM['sources']['inputs']['raw_path'],read_only=True,data_only=True)
    train=np.array([r[1] for r in w['Day-ahead (2021)'].iter_rows(min_row=2,max_row=1441,values_only=True)],float)
    w.close()
    assert len(train)==1440
    return train.reshape(-1,24).mean(axis=0)

def available_plan(d,T,t,observed_outage,elapsed,max_outage,backup,lead=0,lookahead=None):
    n=T-t if lookahead is None else min(lookahead,T-t)
    primary=np.full(n,d.replenish_factor*1.5494*d.qmean)
    bcap=np.zeros(n)
    if observed_outage:
        remaining=min(n,max(0,max_outage-elapsed))
        primary[:remaining]=0
        bcap[:remaining]=backup
        bcap[:max(0,lead-elapsed)]=0
    return primary,bcap

def causal_inputs(d,T,t,state,qprev,published_prices,past_prices,training_forecast,
                  observed_outage,mode,max_outage_h,backup,horizon=72,lead=0,
                  forecast_mode='rolling',emergency=False):
    """Controller inputs: no realized repair date or unrevealed price argument."""
    assert T%24==t%24==max_outage_h%24==horizon%24==0
    assert 0<=t<T and 0<=max_outage_h<=T and horizon>=24 and lead>=0
    assert len(published_prices)==24 and len(past_prices)==t
    h0=min(1/30,.072864*d.hydrogen_volume*65/30*d.usable_hydrogen_fraction)
    if forecast_mode=='rolling' and t:
        forecast=np.asarray(past_prices[-168:]).reshape(-1,24).mean(axis=0)
    else:forecast=training_forecast.copy()
    n=T-t if mode=='guard' or emergency else min(horizon,T-t)
    if mode not in ['guard','mpc','credit_mpc','certificate_mpc']:raise ValueError(mode)
    prices=np.tile(forecast,int(np.ceil(n/24)))[:n]
    prices[:24]=published_prices
    primary,bcap=available_plan(d,T,t,observed_outage,t,max_outage_h,backup,lead=lead,lookahead=n)
    end={'C':d.carbon_initial,'H':h0,'M':d.product_initial} if t+n==T else {'H':h0}
    credit=np.zeros(n);mfloor=np.zeros(n)
    if mode in ['credit_mpc','certificate_mpc'] and observed_outage:
        for k in range(n):
            s=t+k+1
            future_due=max(0,(max_outage_h-s+s%24)//24)*24*d.qmean
            # Lead-time hours with no delivery do not count as reachable backup.
            future_backup=backup*max(0,max_outage_h-max(s,lead))
            credit[k]=max(0,1.5494*future_due-future_backup)
    if mode=='certificate_mpc':
        fp,fb=available_plan(d,T,t,observed_outage,t,max_outage_h,backup,lead=lead)
        for k in range(n):
            s=t+k+1;future_due=((T-s+s%24)//24)*24*d.qmean
            credit[k]=max(credit[k],d.carbon_initial+1.5494*d.product_initial+
                          1.5494*future_due-float((fp[k+1:]+fb[k+1:]).sum()))
            mfloor[k]=max(0,d.product_initial+future_due-d.qmax*(T-s))
    payload=dict(state=state,qprev=qprev,prices=prices.tolist(),primary=primary.tolist(),
                 backup=bcap.tolist(),terminal=end,credit=credit.tolist(),product_floor=mfloor.tolist())
    digest=hashlib.sha256(json.dumps(payload,sort_keys=True).encode()).hexdigest()
    return dict(prices=prices,primary=primary,backup=bcap,terminal=end,credit=credit,
                product_floor=mfloor,input_sha256=digest)

def shift_witness(plan,observed_recovered,primary_cap):
    """Standard retained feasible-tail fallback, qualified interchangeable CO2.

    Price changes affect the objective only. Earlier repair permits replacement
    of backup by primary input, provided recovered gate capacity >= backup cap.
    The remaining q/H/C/M path, ramps, deadlines and final stocks are unchanged.
    """
    candidate=plan.copy()
    if observed_recovered:
        candidate['primary']=candidate.primary+candidate.backup
        candidate['backup']=0.
    assert candidate.primary.max()<=primary_cap+1e-6
    return candidate

def dispatch(actual_prices,d,*,mode,realized_outage_h,backup,max_outage_h=504,
             horizon=72,lead=0,time_limit=20,save_dir=None,forecast_mode='rolling'):
    assert len(actual_prices)%24==0 and realized_outage_h%24==0
    assert 0<=realized_outage_h<=max_outage_h<=len(actual_prices)
    assert horizon>=24 and horizon%24==0 and max_outage_h%24==0 and lead>=0
    T=len(actual_prices);training_forecast=price_forecast()
    h0=min(1/30,.072864*d.hydrogen_volume*65/30*d.usable_hydrogen_fraction)
    state={'C':d.carbon_initial,'H':h0,'M':d.product_initial};qprev=d.qmean
    trace=[];logs=[];status='full_service';emergency=False
    retained_tail=None;admission=None
    if mode=='guard':
        assert d.replenish_factor*1.5494*d.qmean>=backup, 'Shift guarantee requires interchangeable recovered flow >= backup.'
        ap,ab=available_plan(d,T,0,True,0,max_outage_h,backup,lead=lead)
        admission,retained_tail=solve(np.zeros(T),d,0,backup=backup,feasibility=True,
                         primary_caps=ap,backup_caps=ab,time_limit=30)
        if retained_tail is None:raise RuntimeError('No feasible admission witness; no service guarantee is asserted.')
    for t in range(0,T,24):
        # Rolling profile uses only completely revealed PAST days. The first
        # dispatch is initialized from an independent published 2021 input.
        if forecast_mode=='rolling' and t:
            forecast=actual_prices[max(0,t-168):t].reshape(-1,24).mean(axis=0)
        else:forecast=training_forecast.copy()
        # Driver reveals only the current state of the supplier at day start.
        # Test episodes have day-boundary transitions; intra-day hidden changes
        # are outside this evaluated policy's scope.
        observed_outage=t<realized_outage_h
        inputs=causal_inputs(d,T,t,state,qprev,actual_prices[t:t+24],actual_prices[:t],
                            training_forecast,observed_outage,mode,max_outage_h,backup,horizon,lead,
                            forecast_mode,emergency)
        prices,primary,bcap=inputs['prices'],inputs['primary'],inputs['backup'];n=len(prices)
        end,credit,mfloor=inputs['terminal'],inputs['credit'],inputs['product_floor']
        summary,df=solve(prices,d,0,backup=backup,graph_hours=min(24,n),
                         initial=state,initial_load=qprev,terminal=end,
                         primary_caps=primary,backup_caps=bcap,credit_floor=credit,product_floor=mfloor,
                         allow_shortfall=emergency,time_limit=time_limit)
        log=dict(day=t//24,observed_outage=observed_outage,known_price_hours=min(24,n),
                 assumed_remaining_outage_h=max(0,max_outage_h-t) if observed_outage else 0,
                 forecast=('Past 1-7 complete revealed days, 2021 public-profile initialization' if forecast_mode=='rolling'
                           else 'Fixed pretest hour-of-day means of public Mucci2021'),
                 causal_input_sha256=inputs['input_sha256'],
                 terminal_target=end,**summary)
        logs.append(log)
        if mode=='guard' and df is None:
            # Do not turn a missing cost incumbent into a physical service loss.
            # The already feasible tail is independent of unrevealed prices.
            log['cost_optimization_unavailable']=True
            df=shift_witness(retained_tail,not observed_outage,d.replenish_factor*1.5494*d.qmean)
            log['retained_tail_fallback']=dict(source='Feasible admission/reoptimized tail',
                         rows=len(df),input_sha256=inputs['input_sha256'])
        if df is None:
            # Retain the failure, then execute a declared causal emergency
            # service-first fallback. No unexecuted tail is misreported as loss.
            log['hard_service_failed']=True
            emergency=True
            print(json.dumps(dict(mode=mode,failed_day=t//24,status=summary['status'],message=summary['message'])),flush=True)
            n=T-t;prices=np.tile(forecast,int(np.ceil(n/24)))[:n]
            prices[:24]=actual_prices[t:t+24]
            primary,bcap=available_plan(d,T,t,observed_outage,t,max_outage_h,backup,lead=lead)
            summary,df=solve(prices,d,0,backup=backup,graph_hours=24,initial=state,
                             initial_load=qprev,primary_caps=primary,backup_caps=bcap,
                             allow_shortfall=True,time_limit=time_limit)
            log['emergency_fallback']=summary
            log['fallback_input_sha256']=hashlib.sha256(json.dumps(dict(state=state,qprev=qprev,
                   prices=prices.tolist(),primary=primary.tolist(),backup=bcap.tolist(),
                   terminal={'C':d.carbon_initial,'H':h0,'M':d.product_initial},allow_shortfall=True),
                   sort_keys=True).encode()).hexdigest()
            if df is None:
                status='fallback_no_feasible_incumbent';break
        if mode=='guard':retained_tail=df.iloc[24:].copy()
        action=df.iloc[:24].copy()
        action['hour']=np.arange(t,t+24);action['day']=t//24
        action['policy']=mode;action['observed_outage']=observed_outage
        # Forecast-objective values are discarded; executed values use revealed
        # current-day prices only. Physical nonlinear replay is exact.
        action['realized_price_EUR_MWh']=actual_prices[t:t+24]
        action['realized_operating_cost_EUR']=(action.exact_power_MW*action.realized_price_EUR_MWh+
                                              action.exact_cooling_MW+50*(action.primary+action.backup))
        trace.append(action)
        state={s:float(action[s].iloc[-1]) for s in ['C','H','M']}
        qprev=float(action.q.iloc[-1])
        if save_dir:
            save_dir.mkdir(parents=True,exist_ok=True)
            (save_dir/f'{mode}_progress.json').write_text(json.dumps(logs,indent=2),encoding='utf-8')
        if t%168==0:print(json.dumps(dict(mode=mode,day=t//24,C=state['C'],M=state['M'],seconds=summary['seconds'])),flush=True)
    result=pd.concat(trace,ignore_index=True) if trace else None
    due=T*d.qmean
    if result is None:delivered=0
    else:delivered=float(result.delivery_t.sum())
    all_solves=[l for l in logs]+[l['emergency_fallback'] for l in logs if 'emergency_fallback' in l]
    incumbent_solves=[s for s in all_solves if s.get('feasible_incumbent')]
    report=dict(policy=mode,status=status,realized_outage_h=realized_outage_h,
                design_max_outage_h=max_outage_h,backup_t_day=backup*24,complete_hours=0 if result is None else len(result),
                obligation_t=due,delivered_t=delivered,undelivered_or_unexecuted_t=due-delivered,
                measured_model_shortfall_t=0 if result is None else float(result.shortfall.sum()),
                unexecuted_hours=T if result is None else T-len(result),
                hard_service_failure_days=[l['day'] for l in logs if l.get('hard_service_failed')],
                terminal=state,total_exact_operating_cost_EUR=None if result is None else float(result.realized_operating_cost_EUR.sum()),
                worst_solve_seconds=max(l['seconds'] for l in all_solves),
                maximum_solver_constraint_violation=max([l.get('maximum_constraint_violation',0) for l in incumbent_solves],default=None),
                maximum_feasible_incumbent_mip_gap=max([l['mip_gap'] for l in incumbent_solves if l.get('mip_gap') is not None],default=None),
                price_information='Current whole-day prices assumed released before daily dispatch; future forecasts use past revealed prices only.',
                forecast_mode=forecast_mode,
                admission_witness=admission,
                event_information='Only supplier status at each daily boundary; unknown actual future return. Single onset at day zero, day-aligned repair.',
                precollection_capacity_enforced=True,
                optimization=('24h nonlinear-graph PWL cost plus remaining-episode native service tail; tail power is a forecast-cost surrogate.'
                              if mode=='guard' else '72h native service horizon, first24h PWL cost; credit certificates as policy-specific. '
                              'After a failure, latched service-first remaining-episode emergency. Tail power is a forecast-cost surrogate.'))
    if status=='full_service':
        if abs(delivered-due)>1e-5:
            report['status']='executed_with_delivery_deficits'
        assert state['C']>=d.carbon_initial-1e-6 and state['H']>=h0-1e-6 and state['M']>=d.product_initial-1e-6
    if save_dir:
        if result is not None:result.to_csv(save_dir/f'{mode}_trace.csv',index=False)
        (save_dir/f'{mode}_report.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    print(json.dumps(report),flush=True)
    return report,result,logs

def main():
    import argparse
    p=argparse.ArgumentParser();p.add_argument('--days',type=int,default=60)
    p.add_argument('--outage-days',type=int,default=21);p.add_argument('--mode',default='guard',choices=['guard','mpc','credit_mpc','certificate_mpc'])
    p.add_argument('--backup-t-day',type=float,default=242.12);p.add_argument('--segments',type=int,default=4)
    p.add_argument('--price-source',default='mucci',choices=['mucci','DK1'])
    p.add_argument('--forecast',default='rolling',choices=['rolling','fixed']);a=p.parse_args()
    d=Design(segments=a.segments)
    if a.price_source=='mucci':
        w=load_workbook(PARAM['sources']['inputs']['raw_path'],read_only=True,data_only=True)
        prices=np.array([r[1] for r in w['Day-ahead'].iter_rows(min_row=2,max_row=24*a.days+1,values_only=True)],float);w.close()
    else:
        frame=pd.read_csv(ROOT/'results/DK1_hourly_replay_202510_202609.csv')
        # Current data contain UTC timestamps. Start midnight Danish local time
        # on 1Jan2026 and stay before the March DST transition.
        timecol=next(c for c in frame.columns if 'time' in c.lower())
        ts=pd.to_datetime(frame[timecol],utc=True)
        subset=frame.loc[ts>=pd.Timestamp('2025-12-31T23:00:00Z')].iloc[:24*a.days]
        pricecol=next(c for c in frame.columns if 'price' in c.lower())
        prices=subset[pricecol].to_numpy(float)
    assert len(prices)==24*a.days
    out=ROOT/f'results/causal_{a.price_source}_{a.days}d_{a.outage_days}out_{a.mode}_{a.segments}seg_{a.forecast}'
    report,df,logs=dispatch(prices,d,mode=a.mode,realized_outage_h=24*a.outage_days,
                         backup=a.backup_t_day/24,save_dir=out,forecast_mode=a.forecast)
    (out/'design.json').write_text(json.dumps(asdict(d),indent=2),encoding='utf-8')

if __name__=='__main__':main()
