"""Extract public input cell locators and inspect model-output balances, without scaling."""
from pathlib import Path
import json, io, zipfile
import numpy as np
from openpyxl import load_workbook

from runtime import ROOT, RAW, REPO_ROOT
INPUT = ROOT/'cache/giehl_inputs/paper_energy_2024_inputs/paper_energy_2024_input_base_10op.xlsx'
OUTPUT = RAW/'paper_energy_2025_output_raw.zip'

def main():
    report = {'input': str(INPUT), 'output_archive': str(OUTPUT), 'input_cells': {}}
    w = load_workbook(INPUT, read_only=True, data_only=True)
    for name in ['Nodes','Object__node_node','Object__to_from_node','Definition_parameters_duration',
                 'Definition_parameters','Variable_Eff','Operating_points']:
        s=w[name]
        report['input_cells'][name]=[
            {'row':i, 'values': list(r)} for i,r in enumerate(s.values,1)]
    prices=list(w['Energy_prices'].values)
    report['prices']={'n':len(prices)-5,'first':prices[5],'last':prices[-1]}
    w.close()
    with zipfile.ZipFile(OUTPUT) as z:
        name='paper_energy_2025_output_raw/Output_exported_paper_energy_2024_base_10op.xlsx'
        w=load_workbook(io.BytesIO(z.read(name)),read_only=True,data_only=True)
        s=w.worksheets[0]
        rows=s.iter_rows(values_only=True)
        headers=[next(rows) for _ in range(4)]
        keys=[tuple('' if h[j] is None else str(h[j]) for h in headers) for j in range(len(headers[0]))]
        data=np.array([[float(v) if isinstance(v,(int,float)) else np.nan for v in r] for r in rows])
        report['flow_columns']=[]
        for i,k in enumerate(keys):
            if k[0] in ['unit_flow','connection_flow','node_state','units_on']:
                a=data[:,i]
                if any(x in str(k) for x in ['electrolyzer','ch3oh','co2','h2','steam_plant']):
                    report['flow_columns'].append(dict(index=i,key=k,sum=float(np.nansum(a)),
                             max=float(np.nanmax(a)),min=float(np.nanmin(a))))
        def col(k): return data[:,keys.index(k)]
        checks=[('reactor_h2_co2', ('unit_flow','ch3oh_reactor','from_node','h2'),
                 ('unit_flow','ch3oh_reactor','from_node','vaporized_co2'),5.173496722238861),
                ('reactor_h2_raw',('unit_flow','ch3oh_reactor','from_node','h2'),
                 ('unit_flow','ch3oh_reactor','to_node','raw_ch3oh'),1.041666666666667),
                ('dist_raw_product',('unit_flow','dist_tower','from_node','raw_ch3oh'),
                 ('unit_flow','dist_tower','to_node','ch3oh'),1.257861635220126),
                ('vapor_electric_co2',('unit_flow','co2_vaporizer','from_node','power'),
                 ('unit_flow','co2_vaporizer','from_node','co2'),0.003529412)]
        report['checks']=[dict(name=n,max_abs_residual=float(np.nanmax(np.abs(col(a)-r*col(b)))))
                          for n,a,b,r in checks]
        w.close()
    path=ROOT/'results/giehl_input_output_audit.json'
    path.write_text(json.dumps(report,indent=2),encoding='utf-8')
    print(json.dumps(dict(prices=report['prices'],checks=report['checks'],columns=report['flow_columns']),indent=2))

if __name__=='__main__': main()
